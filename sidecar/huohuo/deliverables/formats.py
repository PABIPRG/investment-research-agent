"""纯格式生成；不抓取资料、不执行模型生成的代码或公式。"""
import csv
import io
import json
import re
from datetime import date

from ..contracts import Invalid

MEDIA = {'csv': 'text/csv; charset=utf-8', 'json': 'application/json; charset=utf-8',
         'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'}


def json_bytes(value):
    return (json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + '\n').encode('utf-8')


def cell(value, column):
    if value is None:
        return None
    kind = column['type']
    if kind == 'text':
        if not isinstance(value, str):
            raise Invalid('文本列必须为文本；证券代码不能从数值猜补')
        if len(value) > 16000 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', value):
            raise Invalid('单元格文本过长或含控制字符')
    elif kind in ('number', 'percent'):
        if type(value) not in (int, float) or not (-1e15 < value < 1e15):
            raise Invalid('数值列必须为有限数值或 null')
    elif kind == 'date':
        try:
            if not isinstance(value, str) or len(value) != 10:
                raise ValueError()
            date.fromisoformat(value)
        except ValueError:
            raise Invalid('日期必须为 YYYY-MM-DD 或 null') from None
    return value


def label(column):
    return column['label'] + (f"（{column['unit']}）" if column.get('unit') else '')


def csv_text(value):
    if value is None:
        return ''
    if isinstance(value, str):
        # 单引号是显式文本标记，不借公式保留代码；机器无损消费请用 JSON/XLSX。
        if re.match(r'^[\s\ufeff]*[=+@-]', value) or re.fullmatch(r'0\d+', value) or value.startswith("'"):
            return "'" + value
    return value


def render_file(fmt, tables, metadata):
    if fmt == 'json':
        return json_bytes({'metadata': metadata, 'sources': metadata['sources'], 'tables': tables})
    if fmt == 'csv':
        if len(tables) != 1:
            raise Invalid('CSV 仅支持一个表；多表请选择 Excel 或 JSON')
        table = tables[0]
        stream = io.StringIO(newline='')
        writer = csv.writer(stream, lineterminator='\r\n')
        writer.writerow([csv_text(label(c)) for c in table['columns']] + ['来源证据', '数据性质', '生成时间', '数据限制'])
        for row in table['rows']:
            writer.writerow([csv_text(row.get(c['key'])) for c in table['columns']] + [
                ','.join(metadata['evidence_ids']), metadata['content_kind'],
                metadata['created_at'], csv_text(metadata['limitations'])])
        return stream.getvalue().encode('utf-8-sig')
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    wb = Workbook()
    wb.remove(wb.active)
    for table in tables:
        ws = wb.create_sheet(table['name'])
        ws.append([label(c) for c in table['columns']])
        for row in table['rows']:
            ws.append([None] * len(table['columns']))
            for i, column in enumerate(table['columns'], 1):
                value = row.get(column['key'])
                target = ws.cell(ws.max_row, i)
                target.value = date.fromisoformat(value) if value is not None and column['type'] == 'date' else value
                if column['type'] == 'text' and value is not None:
                    target.data_type = 's'
                target.number_format = {'text': '@', 'percent': '0.00%', 'number': '0.########', 'date': 'yyyy-mm-dd'}[column['type']]
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for i, column in enumerate(table['columns'], 1):
            from openpyxl.utils import get_column_letter
            ws.column_dimensions[get_column_letter(i)].width = min(45, max(14, len(label(column)) * 2 + 4))
        for target in ws[1]:
            target.data_type = 's'
            target.font = Font(bold=True)
            target.alignment = Alignment(wrap_text=True)
    notes = wb.create_sheet('交付说明')
    for key, value in [('标题', metadata['title']), ('内容性质', metadata['content_kind']),
                       ('生成时间', metadata['created_at']), ('来源证据', ', '.join(metadata['evidence_ids'])),
                       ('限制', metadata['limitations']), ('未知值', '空单元格表示未知或未提供，不是零'),
                       ('说明', '来源原始记录及字段单位见配套交付说明 JSON；此文件为生成时快照')]:
        notes.append([key, value])
        notes.cell(notes.max_row, 2).data_type = 's'
    notes.column_dimensions['A'].width = 16
    notes.column_dimensions['B'].width = 90
    for row in notes:
        for target in row:
            target.alignment = Alignment(wrap_text=True, vertical='top')
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()

import csv
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from huohuo.contracts import Invalid
from huohuo.deliverables.service import Deliverables


class DeliverablesTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name) / 'workspace'
        self.evidence = {'status': 'ok', 'data': {'rows': [
            {'code': '000001', 'name': '平安银行', 'price': 12.3, 'cost': None},
        ]}, 'source': 'fixture', 'observed_at': '2026-10-29T08:00:00Z'}
        self.reads = []

        def read(principal, session, eid):
            self.reads.append((principal, session, eid))
            if (principal, session, eid) != ('alice', 'chat', 'ev1'):
                raise Invalid('证据不属于当前对话')
            return self.evidence

        self.service = Deliverables(self.root, read)

    def tearDown(self):
        self.tmp.cleanup()

    def spec(self, fmt='xlsx'):
        return {'format': fmt, 'title': '银行研究', 'content_kind': 'evidence',
                'evidence_ids': ['ev1'], 'limitations': '合成测试，非当前行情', 'tables': [
                    {'name': '数据', 'evidence_id': 'ev1', 'records_path': ['data', 'rows'],
                     'columns': [{'key': 'code', 'label': '代码', 'type': 'text'},
                                 {'key': 'name', 'label': '名称', 'type': 'text'},
                                 {'key': 'price', 'label': '价格', 'type': 'number', 'unit': '元'},
                                 {'key': 'cost', 'label': '成本', 'type': 'number', 'unit': '元'}]}]}

    def test_source_values_and_xlsx_types(self):
        from openpyxl import load_workbook
        result = self.service.create('alice', 'chat', 'call1', self.spec())
        meta, raw = self.service.download('alice', result['id'])
        wb = load_workbook(io.BytesIO(raw))
        ws = wb['数据']
        self.assertEqual([ws.cell(2, i).value for i in range(1, 5)], ['000001', '平安银行', 12.3, None])
        self.assertEqual(ws['A2'].data_type, 's')
        self.assertEqual(ws.freeze_panes, 'A2')
        self.assertEqual(ws.auto_filter.ref, 'A1:D2')
        self.assertIn('交付说明', wb.sheetnames)
        self.assertEqual(self.reads, [('alice', 'chat', 'ev1')])
        self.assertEqual(meta['sha256'], result['sha256'])

    def test_json_preserves_unknown_and_nested_source(self):
        s = self.spec('json'); s.pop('tables')
        result = self.service.create('alice', 'chat', 'json', s)
        _, raw = self.service.download('alice', result['id'])
        data = json.loads(raw)
        self.assertEqual(data['sources']['ev1'], self.evidence)
        self.assertIsNone(data['sources']['ev1']['data']['rows'][0]['cost'])
        self.assertEqual(data['metadata']['content_kind'], 'evidence')

    def test_csv_encoding_injection_and_code(self):
        self.evidence['data']['rows'][0]['name'] = '=HYPERLINK("x")'
        result = self.service.create('alice', 'chat', 'csv', self.spec('csv'))
        _, raw = self.service.download('alice', result['id'])
        self.assertTrue(raw.startswith(b'\xef\xbb\xbf'))
        rows = list(csv.reader(io.StringIO(raw.decode('utf-8-sig'))))
        self.assertEqual(rows[1][:4], ["'000001", '\'=HYPERLINK("x")', '12.3', ''])
        self.assertIn('来源证据', rows[0])

    def test_replay_restart_collision_and_ownership(self):
        first = self.service.create('alice', 'chat', 'same', self.spec())
        restarted = Deliverables(self.root, lambda *_: self.fail('重放不得重新获取数据'))
        self.assertEqual(restarted.create('alice', 'chat', 'same', self.spec()), first)
        self.assertEqual(restarted.inspect('alice', 'chat', first['id']), first)
        with self.assertRaises(Invalid):
            restarted.create('alice', 'chat', 'same', {**self.spec(), 'title': '不同'})
        for user, session in [('bob', 'chat'), ('alice', 'other')]:
            with self.assertRaises(Invalid): restarted.inspect(user, session, first['id'])
        with self.assertRaises(Invalid): restarted.download('bob', first['id'])
        with self.assertRaises(Invalid): restarted.download('alice', '../business.sqlite')
        other = self.service.create('alice', 'chat', 'different', self.spec())
        self.assertNotEqual(other['id'], first['id'])

    def test_reject_forged_rows_sources_types_and_paths(self):
        cases = []
        forged = self.spec(); forged['tables'][0]['rows'] = [{'code': '999999'}]; cases.append(forged)
        foreign = self.spec(); foreign['evidence_ids'] = ['foreign']; cases.append(foreign)
        cases.append({**self.spec(), 'path': '/tmp/escape'})
        cases.append({**self.spec(), 'title': '../escape'})
        for i, spec in enumerate(cases):
            with self.assertRaises(Invalid): self.service.create('alice', 'chat', str(i), spec)
        self.evidence['data']['rows'][0]['price'] = 'unknown'
        with self.assertRaises(Invalid): self.service.create('alice', 'chat', 'bad-number', self.spec())

    def test_atomic_failure_and_file_missing_or_modified(self):
        with patch('huohuo.deliverables.service.render_file', side_effect=OSError('disk full')):
            with self.assertRaises(OSError): self.service.create('alice', 'chat', 'fail', self.spec())
        self.assertEqual(list(self.root.rglob('manifest.json')), [])
        result = self.service.create('alice', 'chat', 'good', self.spec())
        path = Path(result['file_path'])
        path.write_bytes(b'changed')
        with self.assertRaises(Invalid): self.service.download('alice', result['id'])
        path.unlink()
        with self.assertRaises(Invalid): self.service.download('alice', result['id'])

    def test_model_table_is_labelled_and_formula_is_literal(self):
        from openpyxl import load_workbook
        spec = {'format': 'xlsx', 'title': '分析摘要', 'content_kind': 'model_analysis',
                'evidence_ids': ['ev1'], 'limitations': '模型整理，需核对', 'tables': [
                    {'name': '摘要', 'columns': [{'key': 'text', 'label': '判断', 'type': 'text'}],
                     'rows': [{'text': '=1+1'}, {'text': None}]}]}
        r = self.service.create('alice', 'chat', 'analysis', spec)
        _, raw = self.service.download('alice', r['id'])
        ws = load_workbook(io.BytesIO(raw))['摘要']
        self.assertEqual(ws['A2'].value, '=1+1')
        self.assertEqual(ws['A2'].data_type, 's')
        meta, raw = self.service.download('alice', r['id'], 'metadata')
        self.assertEqual(json.loads(raw)['content_kind'], 'model_analysis')
        self.assertEqual(meta['filename'], '交付说明.json')

    def test_empty_multisheet_headers_types_and_limits(self):
        from openpyxl import load_workbook
        self.evidence['status'] = 'empty'
        self.evidence['data']['rows'] = []
        spec = self.spec()
        second = {**spec['tables'][0], 'name': '空结果'}
        spec['tables'].append(second)
        spec['tables'][0]['columns'][0]['label'] = '=1+1'
        result = self.service.create('alice', 'chat', 'empty', spec)
        _, raw = self.service.download('alice', result['id'])
        wb = load_workbook(io.BytesIO(raw))
        self.assertEqual(wb.sheetnames, ['数据', '空结果', '交付说明'])
        self.assertEqual(wb['数据']['A1'].data_type, 's')
        self.assertEqual(wb['数据'].max_row, 1)
        with self.assertRaises(Invalid):
            self.service.create('alice', 'chat', 'two-csv', {**spec, 'format': 'csv'})
        spec['tables'] = [second] * 7
        with self.assertRaises(Invalid): self.service.create('alice', 'chat', 'seven', spec)

    def test_staging_failure_symlink_and_concurrent_replay(self):
        from concurrent.futures import ThreadPoolExecutor
        with patch('huohuo.deliverables.service.os.fsync', side_effect=OSError('disk full')):
            with self.assertRaises(OSError): self.service.create('alice', 'chat', 'write-fail', self.spec())
        self.assertEqual(list(self.root.rglob('manifest.json')), [])
        self.assertEqual(list((self.root / '.huohuo-deliverables').iterdir()), [])
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(lambda _: self.service.create('alice', 'chat', 'race', self.spec()), range(2)))
        self.assertEqual(results[0], results[1])
        self.assertEqual(self.service.by_call('alice', 'chat', 'race'), results[0])
        path = Path(results[0]['file_path'])
        outside = Path(self.tmp.name) / 'outside'
        outside.write_bytes(b'private')
        path.unlink(); path.symlink_to(outside)
        with self.assertRaises(Invalid): self.service.download('alice', results[0]['id'])

    def test_published_but_unacknowledged_is_reconcilable_and_metadata_is_required(self):
        import os
        if os.name != 'posix':
            self.skipTest('仅 POSIX 执行目录 fsync 故障注入')
        sync = os.fsync
        calls = []
        def uncertain(descriptor):
            calls.append(descriptor)
            if len(calls) == 4:
                raise OSError('directory acknowledgement lost')
            sync(descriptor)
        with patch('huohuo.deliverables.service.os.fsync', side_effect=uncertain):
            with self.assertRaises(OSError): self.service.create('alice', 'chat', 'uncertain', self.spec())
        result = self.service.by_call('alice', 'chat', 'uncertain')
        Path(result['metadata_path']).unlink()
        with self.assertRaises(Invalid): self.service.create('alice', 'chat', 'uncertain', self.spec())


if __name__ == '__main__':
    unittest.main()
